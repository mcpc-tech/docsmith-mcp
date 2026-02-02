"""
Excel document handler - read/write Excel files
"""
import json
import sys
from pathlib import Path

def read_excel(file_path: str, sheet_name: str = None, page: int = None, page_size: int = 100):
    """Read Excel file with optional pagination"""
    import openpyxl
    
    wb = openpyxl.load_workbook(file_path, data_only=True)
    
    # Handle sheet_name selection
    if sheet_name is None:
        # Use first sheet if not specified
        sheet_name = wb.sheetnames[0]
    elif sheet_name not in wb.sheetnames:
        # If sheet_name not found, try to interpret as 1-based index
        try:
            sheet_index = int(sheet_name) - 1
            if 0 <= sheet_index < len(wb.sheetnames):
                sheet_name = wb.sheetnames[sheet_index]
            else:
                # Index out of range, use first sheet
                sheet_name = wb.sheetnames[0]
        except (ValueError, IndexError):
            # Not a valid number or other error, use first sheet
            sheet_name = wb.sheetnames[0]
    
    ws = wb[sheet_name]
    
    # Get all data
    data = []
    for row in ws.iter_rows(values_only=True):
        data.append(row)
    
    total_rows_data = len(data)
    
    # Handle pagination
    if page is not None:
        total_pages = (total_rows_data + page_size - 1) // page_size if total_rows_data > 0 else 1
        start = (page - 1) * page_size
        end = start + page_size
        data = data[start:end]
    else:
        total_pages = 1
    
    return {
        "sheet_name": sheet_name,
        "sheets": wb.sheetnames,
        "total_rows": ws.max_row,
        "total_cols": ws.max_column,
        "current_page": page,
        "page_size": page_size if page else None,
        "total_pages": total_pages,
        "data": data
    }

def get_excel_info(file_path: str):
    """Get Excel file metadata"""
    import openpyxl
    
    wb = openpyxl.load_workbook(file_path, data_only=True)
    info = {
        "sheets": [],
        "file_size": Path(file_path).stat().st_size
    }
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        info["sheets"].append({
            "name": sheet_name,
            "rows": ws.max_row,
            "cols": ws.max_column
        })
    
    return info

def write_excel(file_path: str, data: list, sheet_name: str = "Sheet1"):
    """Write data to Excel file"""
    import openpyxl
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    
    for row in data:
        ws.append(row)
    
    wb.save(file_path)
    return {"success": True, "file_path": file_path}

if __name__ == "__main__":
    command = sys.argv[1]
    file_path = sys.argv[2]
    
    if command == "read":
        sheet = sys.argv[3] if len(sys.argv) > 3 and sys.argv[3] else None
        page = int(sys.argv[4]) if len(sys.argv) > 4 and sys.argv[4] else None
        page_size = int(sys.argv[5]) if len(sys.argv) > 5 else 100
        result = read_excel(file_path, sheet, page, page_size)
    elif command == "info":
        result = get_excel_info(file_path)
    elif command == "write":
        # Data passed as JSON string
        data = json.loads(sys.argv[3])
        sheet = sys.argv[4] if len(sys.argv) > 4 else "Sheet1"
        result = write_excel(file_path, data, sheet)
    else:
        result = {"error": f"Unknown command: {command}"}
    
    print(json.dumps(result, default=str))
