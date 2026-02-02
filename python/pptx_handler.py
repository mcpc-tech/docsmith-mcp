"""
PowerPoint document handler - read/write PPTX files
"""
import json
import sys
from pathlib import Path


def read_pptx(file_path: str, page: int = None, page_size: int = 100):
    """Read PowerPoint presentation with optional pagination by slides"""
    from pptx import Presentation
    
    prs = Presentation(file_path)
    
    # Extract all slides content
    slides = []
    for i, slide in enumerate(prs.slides):
        slide_data = {
            "slide_number": i + 1,
            "title": "",
            "content": [],
            "notes": ""
        }
        
        # Extract text from shapes
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                # Try to detect if it's a title
                if hasattr(shape, "is_placeholder") and shape.is_placeholder:
                    placeholder = shape.placeholder_format
                    if placeholder.type == 1:  # Title placeholder
                        slide_data["title"] = shape.text
                        continue
                
                slide_data["content"].append(shape.text)
            
            # Extract table data
            if hasattr(shape, "table"):
                table_data = []
                for row in shape.table.rows:
                    row_data = [cell.text for cell in row.cells]
                    table_data.append(row_data)
                slide_data["content"].append({"table": table_data})
        
        # Extract notes
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame:
                slide_data["notes"] = notes_frame.text
        
        slides.append(slide_data)
    
    total_slides = len(slides)
    
    # Handle pagination
    if page is not None:
        start = (page - 1) * page_size
        end = start + page_size
        slides = slides[start:end]
        total_pages = (total_slides + page_size - 1) // page_size if total_slides else 1
    else:
        total_pages = 1
    
    return {
        "total_slides": total_slides,
        "slides": slides,
        "current_page": page,
        "page_size": page_size if page else None,
        "total_pages": total_pages
    }


def get_pptx_info(file_path: str):
    """Get PowerPoint metadata"""
    from pptx import Presentation
    
    prs = Presentation(file_path)
    
    info = {
        "slides": len(prs.slides),
        "file_size": Path(file_path).stat().st_size
    }
    
    # Try to get presentation properties
    if prs.core_properties:
        props = prs.core_properties
        metadata = {}
        if props.title:
            metadata["title"] = props.title
        if props.author:
            metadata["author"] = props.author
        if props.subject:
            metadata["subject"] = props.subject
        if props.created:
            metadata["created"] = str(props.created)
        if props.modified:
            metadata["modified"] = str(props.modified)
        
        if metadata:
            info["metadata"] = metadata
    
    return info


def write_pptx(file_path: str, slides_data: list):
    """Write data to PowerPoint presentation"""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    for slide_info in slides_data:
        # Add blank slide
        blank_layout = prs.slide_layouts[6]  # Blank layout
        slide = prs.slides.add_slide(blank_layout)
        
        # Add title if provided
        title = slide_info.get("title", "")
        if title:
            left = Inches(0.5)
            top = Inches(0.5)
            width = Inches(9)
            height = Inches(1)
            title_box = slide.shapes.add_textbox(left, top, width, height)
            title_frame = title_box.text_frame
            title_frame.text = title
            title_frame.paragraphs[0].font.size = Pt(32)
            title_frame.paragraphs[0].font.bold = True
        
        # Add content
        content = slide_info.get("content", [])
        if content:
            left = Inches(0.5)
            top = Inches(2)
            width = Inches(9)
            height = Inches(5)
            content_box = slide.shapes.add_textbox(left, top, width, height)
            text_frame = content_box.text_frame
            
            for item in content:
                if isinstance(item, str):
                    p = text_frame.add_paragraph()
                    p.text = item
                    p.level = 0
    
    prs.save(file_path)
    return {"success": True, "file_path": file_path}


if __name__ == "__main__":
    command = sys.argv[1]
    file_path = sys.argv[2]
    
    if command == "read":
        page = int(sys.argv[3]) if len(sys.argv) > 3 else None
        page_size = int(sys.argv[4]) if len(sys.argv) > 4 else 100
        result = read_pptx(file_path, page, page_size)
    elif command == "info":
        result = get_pptx_info(file_path)
    elif command == "write":
        # Data passed as JSON string
        slides_data = json.loads(sys.argv[3])
        result = write_pptx(file_path, slides_data)
    else:
        result = {"error": f"Unknown command: {command}"}
    
    print(json.dumps(result, default=str))
